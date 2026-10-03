#!/usr/bin/env python3
"""Build publication-facing SRA metadata and FASTQ manifests from the checked inventory."""
import argparse
import csv
import json
from collections import defaultdict
from pathlib import Path

from public_sample_names import PUBLIC_SAMPLES, public_name

HEADERS = ["sample_name", "library_ID", "title", "library_strategy", "library_source",
           "library_selection", "library_layout", "platform", "instrument_model",
           "design_description", "filetype", "filename", "filename2", "filename3",
           "filename4", "assembly", "fasta_file"]
LIBRARY_METHOD = ("RNA underwent bead-based poly(A) selection, cDNA synthesis, and PCR amplification "
          "with the VAHTS Universal V10 RNA-seq Library Prep Kit for Illumina "
          "(Vazyme #NR606-02). Equimolar libraries were sequenced as 150-bp paired-end "
          "reads on an Illumina NovaSeq X Plus.")
DESIGN_TGN = "Total RNA was extracted from TGN cultures using the RNeasy kit (Qiagen #74104). " + LIBRARY_METHOD
DESIGN_HDRG = "Pooled human DRG RNA comparator material was supplied as RNA. " + LIBRARY_METHOD


def filename(item):
    name = public_name(item["sample"])
    cohort = item["sample"].split("_", 1)[0]
    if name.startswith("Cheng_hDRG"):
        return f"{name}_{cohort}_R{item['mate']}.fastq.gz"
    return f"{name}_R{item['mate']}.fastq.gz"


def title(public):
    if public.startswith("Cheng_hDRG"):
        return f"Pooled human dorsal root ganglion RNA comparator {public[-1]}; two technical sequencing runs"
    if public == "Cheng_TGN_Pilot_W1":
        return "hiPSC-derived trigeminal sensory neuron pilot culture, week 1"
    lot, week = public.rsplit("_", 2)[-2:]
    return f"hiPSC-derived trigeminal sensory neuron culture, lot {lot[1:]}, week {week[1:]}"


def write_tsv(path, rows, fields):
    with path.open("w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=fields, delimiter="\t", lineterminator="\n")
        writer.writeheader()
        writer.writerows(rows)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--raw-plan", type=Path, required=True)
    parser.add_argument("--hashes", type=Path)
    parser.add_argument("--outdir", type=Path, required=True)
    parser.add_argument("--template", type=Path, required=True)
    parser.add_argument("--dest-prefix", required=True)
    args = parser.parse_args()
    args.outdir.mkdir(parents=True, exist_ok=True)
    plan = json.loads(args.raw_plan.read_text())
    objects = plan["objects"]
    by_sample = defaultdict(list)
    for item in objects:
        enriched = dict(item, public_sample=public_name(item["sample"]), filename=filename(item))
        by_sample[enriched["public_sample"]].append(enriched)
    if set(by_sample) != set(PUBLIC_SAMPLES.values()) or len(objects) != 34:
        raise ValueError("Expected exactly 15 biological samples and 34 FASTQs")
    if len({filename(item) for item in objects}) != 34:
        raise ValueError("Nonunique public FASTQ filename")
    rows = []
    final_plan = []
    for public in PUBLIC_SAMPLES.values():
        items = sorted(by_sample[public], key=lambda x: (x["sample"].split("_", 1)[0], x["mate"]))
        if len(items) != (4 if public.startswith("Cheng_hDRG") else 2):
            raise ValueError(f"Wrong FASTQ count for {public}")
        if [x["mate"] for x in items] != ([1, 2, 1, 2] if len(items) == 4 else [1, 2]):
            raise ValueError(f"Unpaired mates for {public}")
        names = [x["filename"] for x in items]
        rows.append(dict(zip(HEADERS, [public, f"{public}_RNA", title(public), "RNA-Seq",
                                       "TRANSCRIPTOMIC", "PolyA", "paired", "ILLUMINA",
                                       "Illumina NovaSeq X Plus",
                                       DESIGN_HDRG if public.startswith("Cheng_hDRG") else DESIGN_TGN, "fastq",
                                       *(names + [""] * (4 - len(names))), "", ""])))
        for item in items:
            final_plan.append({"source_key": item["source_key"],
                               "dest_key": f"{args.dest_prefix.rstrip('/')}/{item['filename']}",
                               "sample_name": public, "filename": item["filename"],
                               "cohort": item["sample"].split("_", 1)[0],
                               "read": item["mate"], "size_bytes": item["size_bytes"]})
    write_tsv(args.outdir / "Cheng_TGN_SRA_metadata.tsv", rows, HEADERS)
    (args.outdir / "private_s3_copy_plan.json").write_text(json.dumps({"bucket": plan["bucket"], "objects": final_plan}, indent=2) + "\n")
    from openpyxl import load_workbook
    book = load_workbook(args.template)
    sheet = book["SRA_data"]
    if [sheet.cell(1, col).value for col in range(1, 18)] != HEADERS:
        raise ValueError("SRA template header changed")
    for row_num, row in enumerate(rows, 2):
        for col_num, field in enumerate(HEADERS, 1):
            sheet.cell(row_num, col_num).value = row[field] or None
    book.save(args.outdir / "Cheng_TGN_SRA_metadata.xlsx")
    if args.hashes:
        with args.hashes.open(newline="") as f:
            hashes = {r["source_key"]: r for r in csv.DictReader(f, delimiter="\t")}
        if set(hashes) != {x["source_key"] for x in final_plan}:
            raise ValueError("Incomplete or extra FASTQ digest records")
        manifest = []
        for item in final_plan:
            digest = hashes[item["source_key"]]
            if int(digest["bytes"]) != item["size_bytes"]:
                raise ValueError(f"Size mismatch for {item['filename']}")
            manifest.append({"sample_name": item["sample_name"], "library_ID": f"{item['sample_name']}_RNA",
                             "filename": item["filename"], "cohort": item["cohort"],
                             "read": item["read"], "bytes": item["size_bytes"],
                             "md5": digest["md5"], "sha256": digest["sha256"]})
        write_tsv(args.outdir / "Cheng_TGN_FASTQ_checksums.tsv", manifest,
                  ["sample_name", "library_ID", "filename", "cohort", "read", "bytes", "md5", "sha256"])
    print(f"15 SRA rows and {len(final_plan)} unique FASTQ filenames validated")


if __name__ == "__main__":
    main()
