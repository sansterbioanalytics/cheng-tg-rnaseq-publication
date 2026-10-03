#!/usr/bin/env python3
"""Reconcile the Human 1.0 BioSample draft with public SRA sample identifiers."""
import argparse
import csv
from pathlib import Path

from public_sample_names import public_name

EXTRA = ["differentiation_lot", "culture_week", "sample_role"]
MISSING = "not collected"


def write_tsv(path, rows, fields):
    with path.open("w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=fields, delimiter="\t", lineterminator="\n")
        writer.writeheader()
        writer.writerows(rows)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", type=Path, required=True)
    parser.add_argument("--sra", type=Path, required=True)
    parser.add_argument("--template", type=Path, required=True)
    parser.add_argument("--outdir", type=Path, required=True)
    args = parser.parse_args()
    args.outdir.mkdir(parents=True, exist_ok=True)
    with args.source.open(newline="") as f:
        source = list(csv.DictReader(f, delimiter="\t"))
    with args.sra.open(newline="") as f:
        sra = list(csv.DictReader(f, delimiter="\t"))
    sra_by_name = {r["sample_name"]: r for r in sra}
    source_by_name = {}
    for row in source:
        old = row["*sample_name"]
        public = {"hDRG_1": "Cheng_hDRG1", "hDRG_2": "Cheng_hDRG2"}.get(old)
        if public is None:
            public = public_name(old)
        if public in source_by_name:
            raise ValueError(f"Duplicate BioSample alias {public}")
        source_by_name[public] = row
    if set(source_by_name) != set(sra_by_name) or len(source_by_name) != 15:
        raise ValueError("BioSample/SRA sample sets do not match")
    fields = list(source[0]) + EXTRA
    rows = []
    for public in sra_by_name:
        row = dict(source_by_name[public])
        row["*sample_name"] = public
        row["sample_title"] = sra_by_name[public]["title"]
        for field in ("*age", "*biomaterial_provider", "*collection_date", "*geo_loc_name", "*sex"):
            if not row[field].strip():
                row[field] = MISSING
        if public.startswith("Cheng_hDRG"):
            comparator = public[-1]
            row["*isolate"] = f"Primary Human DRG {comparator} pooled RNA"
            row["*biomaterial_provider"] = "Takara Bio"
            row["*tissue"] = "dorsal root ganglion"
            if comparator == "1":
                row["*sex"] = "not applicable"
            row["sample_type"] = "tissue sample"
            row["description"] = (
                "Commercial pooled human dorsal root ganglion RNA from Takara Bio; source pool 1 comprises 21 donors; C1 and C2 are technical sequencing runs."
                if comparator == "1" else
                "Second, separate pool of commercial human dorsal root ganglion RNA from Takara Bio; C1 and C2 are technical sequencing runs of this pool."
            )
            row["differentiation_lot"] = "not applicable"
            row["culture_week"] = "not applicable"
            row["sample_role"] = f"hDRG comparator {comparator}"
        else:
            row["*tissue"] = "not applicable"
            row["sample_type"] = "cell culture"
            row["cell_type"] = "hiPSC-derived trigeminal sensory neuron"
            row["cell_subtype"] = "trigeminal sensory neuron"
            tokens = public.split("_")
            if "Pilot" in tokens:
                lot, week = "pilot", "1"
            else:
                lot, week = tokens[-2][1:], tokens[-1][1:]
            row["differentiation_lot"] = lot
            row["culture_week"] = week
            row["sample_role"] = "TGN culture"
            row["description"] = f"hiPSC-derived trigeminal sensory neuron culture; differentiation lot {lot}; week {week}."
        rows.append(row)
    write_tsv(args.outdir / "Cheng_TGN_BioSample_Human1.0.tsv", rows, fields)
    from openpyxl import load_workbook
    book = load_workbook(args.template)
    sheet = book["Human.1.0"]
    for col, field in enumerate(fields, 1):
        sheet.cell(12, col).value = field
    for row_num in range(13, sheet.max_row + 1):
        for col_num in range(1, len(fields) + 1):
            sheet.cell(row_num, col_num).value = None
    for row_num, row in enumerate(rows, 13):
        for col_num, field in enumerate(fields, 1):
            sheet.cell(row_num, col_num).value = row[field] or None
    book.save(args.outdir / "Cheng_TGN_BioSample_Human1.0.xlsx")
    mandatory = [field for field in fields if field.startswith("*")]
    if any(not row[field] for row in rows for field in mandatory):
        raise ValueError("A required Human 1.0 attribute is blank")
    print("15 BioSample rows match SRA sample_name; required fields populated")


if __name__ == "__main__":
    main()
